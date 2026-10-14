"""Phase 4 robustness: config, mixed real-world files, and the always-current index.

Every test here exists because of a real failure found while building Phase 4,
or because the behaviour must hold for documents KORA has never seen.
"""

import json
import os
import time
from pathlib import Path

import pytest

from src.extraction.extractor import ExtractedDocument
from src.ingestion.file_intake import FileInfo, FileType
from src.metadata.document_metadata import (
    apply_fallback_identity, derived_document_id, header_block, identify_document_metadata,
)
from src.retrieval.config import ConfigError, KoraConfig, load_config, parse_config
from src.retrieval.router import Router


# ---------------- config ----------------

def test_missing_config_file_means_safe_defaults(tmp_path):
    cfg = load_config(tmp_path / "nope.json")
    assert cfg.relocation is None and cfg.communication is None
    assert cfg.t("top_k") == 5
    assert any("no config file" in w for w in cfg.warnings)


def test_invalid_config_lists_every_problem():
    with pytest.raises(ConfigError) as e:
        parse_config({"thresholds": {"top_k": 0, "max_lines": 2.5, "nope": 1},
                      "relocation": {"name": "x", "intent_patterns": ["("]}, "surprise": 1})
    message = str(e.value)
    for expected in ("top_k", "max_lines", "nope", "invalid regex", "surprise"):
        assert expected in message


def test_reserved_slots_must_leave_room_for_the_router():
    with pytest.raises(ConfigError):
        parse_config({"thresholds": {"top_k": 3, "global_reserved_slots": 3}})


def test_config_reference_to_missing_document_switches_that_part_off():
    cfg = parse_config({"communication": {"intent_patterns": ["\\btone\\b"], "profiles": {"mars": "DOC-NOPE"},
                                          "country_words": {"mars": "mars"}}})
    cfg.check_references(document_ids={"DOC-A"}, topic_ids=set())
    assert cfg.communication.profiles == {}
    assert any("DOC-NOPE" in w for w in cfg.warnings)


def test_nexora_config_file_is_valid():
    cfg = load_config()
    assert cfg.company == "Nexora Systems" and cfg.relocation and cfg.communication


def test_company_features_are_off_without_config(intel):
    """No config: no relocation/communication intents -- the generic router still works."""
    router = Router.from_outputs(intel.outputs, KoraConfig())
    plan = router.route("What should a new employee moving to Germany prepare during their first week?")
    assert "relocation" not in plan.intents and "communication" not in plan.intents
    assert router.route("Who approves purchases above 10000 euros?").intents == ["approval"]


# ---------------- Phase 2: mixed real-world files ----------------

def _doc(text, name="x.md"):
    return ExtractedDocument(FileInfo(name, FileType.MARKDOWN, len(text), name), text)


def test_metadata_only_read_from_the_header_block():
    """Bug found in Phase 4: a body line 'Status: Archived' marked a whole
    real-world document as archived, and KORA would have hidden it."""
    text = "# IT Policy\n\nIT policy text.\n\nTicket example:\nStatus: Archived\nID: 12345\n"
    m = identify_document_metadata(_doc(text))
    assert m.status == "unknown" and m.document_id == "unknown"
    assert header_block(text) == ""


def test_nexora_style_header_still_read():
    text = "**Document ID:** DOC-1\n**Status:** Active\n\n---\n\n# Title\nBody. Status: Archived\n"
    m = identify_document_metadata(_doc(text))
    assert (m.document_id, m.status, m.id_source) == ("DOC-1", "Active", "header")


def test_documents_without_header_get_labelled_unique_ids():
    """Bug found in Phase 4: every headerless document got the ID 'unknown',
    so unrelated files looked like one document to retrieval."""
    a = apply_fallback_identity(identify_document_metadata(_doc("HR Policy\n\nText.", "policy.md")),
                                _doc("HR Policy\n\nText.", "policy.md"), "hr/policy.md")
    b = apply_fallback_identity(identify_document_metadata(_doc("# IT Policy\n\nText.", "policy.md")),
                                _doc("# IT Policy\n\nText.", "policy.md"), "it/policy.md")
    assert (a.document_id, b.document_id) == ("FILE-HR-POLICY", "FILE-IT-POLICY")
    assert a.id_source == b.id_source == "derived_from_path"
    assert (a.title, b.title) == ("HR Policy", "IT Policy")
    assert a.title_source == "derived_from_text"


def test_derived_id_is_stable_and_path_based():
    assert derived_document_id("Finance/Travel Policy (2024).pdf") == "FILE-FINANCE-TRAVEL-POLICY-2024"


def _write_company(raw: Path):
    """A small documents-only company: no headers, no CSV tables, mixed formats."""
    from docx import Document
    (raw / "hr").mkdir(parents=True)
    (raw / "it").mkdir()
    (raw / "hr" / "policy.md").write_text(
        "# Leave Policy\n\nEmployees get 25 days of annual leave per year. Leave requests longer than "
        "five days need manager approval through the HR portal.\n", encoding="utf-8")
    (raw / "it" / "policy.md").write_text(
        "# Laptop Policy\n\nEvery employee receives a company laptop on the first day. Lost laptops "
        "must be reported to the IT service desk within 24 hours.\n\nTicket example:\nStatus: Archived\n",
        encoding="utf-8")
    d = Document()
    d.add_heading("Remote Work Policy", 1)
    d.add_paragraph("Staff may work remotely up to three days per week with manager agreement. "
                    "Remote work from another country requires HR approval in advance.")
    d.save(raw / "remote.docx")


def test_documents_only_company_builds_and_answers(tmp_path):
    """No CSV tables and no headers: Phase 2 and 3 must complete (they used to
    crash) and KORA must answer with honest metadata notes."""
    from src.index import Kora, Paths
    raw = tmp_path / "raw"
    _write_company(raw)
    paths = Paths(raw_dir=raw, processed_dir=tmp_path / "processed", chunks_dir=tmp_path / "chunks",
                  intelligence_dir=tmp_path / "intelligence", context_dir=tmp_path / "context",
                  manifest=tmp_path / "manifest.json", config=tmp_path / "no_config.json", log_dir=None)
    kora = Kora.open(paths)
    assert kora.last_build["phase3_validation"]["failed"] == 0
    a = kora.ask("How many days of annual leave do employees get?")
    assert a.status == "answered"
    assert a.citations[0].document_id == "FILE-HR-POLICY"
    assert any("does not state its status or version" in n for n in a.notes)
    assert any("derived from its file path" in n for n in a.notes)
    laptop = kora.ask("When must a lost laptop be reported?")
    assert laptop.citations[0].status == "unknown"          # body "Status: Archived" ignored
    assert kora.ask("What is the policy on bringing pets to work?").status == "not_found"


# ---------------- index: always current, never half-built ----------------

@pytest.fixture
def live(tmp_path):
    from src.index import Paths
    raw = tmp_path / "raw"
    _write_company(raw)
    return Paths(raw_dir=raw, processed_dir=tmp_path / "processed", chunks_dir=tmp_path / "chunks",
                 intelligence_dir=tmp_path / "intelligence", context_dir=tmp_path / "context",
                 manifest=tmp_path / "manifest.json", config=tmp_path / "no_config.json",
                 log_dir=tmp_path / "logs")


def test_new_file_is_answerable_on_the_next_question(live):
    from src.index import Kora, check
    kora = Kora.open(live)
    assert kora.ask("What is the bike leasing policy?").status == "not_found"
    (live.raw_dir / "bike.md").write_text("# Bike Leasing\n\nEmployees can lease one bike through payroll "
                                          "for 36 months with HR approval.\n", encoding="utf-8")
    assert check(live).added == ["bike.md"]
    a = kora.ask("What is the bike leasing policy?")
    assert a.status == "answered" and a.citations[0].source_filename == "bike.md"
    assert check(live).current


def test_changed_and_deleted_files_are_detected(live):
    from src.index import Kora, check
    Kora.open(live)
    policy = live.raw_dir / "hr" / "policy.md"
    policy.write_text(policy.read_text(encoding="utf-8").replace("25 days", "30 days"), encoding="utf-8")
    assert check(live).changed == ["hr/policy.md"]
    kora = Kora.open(live)
    assert "30 days" in kora.ask("How many days of annual leave do employees get?").lines[0].text
    (live.raw_dir / "remote.docx").unlink()
    assert check(live).removed == ["remote.docx"]


def test_touching_a_file_without_changing_it_is_not_a_change(live):
    from src.index import Kora, check
    Kora.open(live)
    path = live.raw_dir / "hr" / "policy.md"
    later = time.time() + 5
    os.utime(path, (later, later))
    assert check(live).current


def test_failed_rebuild_keeps_the_previous_index(live, monkeypatch):
    from src import index
    index.Kora.open(live)
    before = (live.chunks_dir / "chunks.jsonl").read_text(encoding="utf-8")
    manifest_before = live.manifest.read_text(encoding="utf-8")
    (live.raw_dir / "new.md").write_text("# New\n\nSomething new to index here.\n", encoding="utf-8")

    def boom(*a, **k):
        raise RuntimeError("simulated Phase 3 failure")
    monkeypatch.setattr("src.build_intelligence.run_intelligence", boom)
    with pytest.raises(index.IndexBuildError, match="previous index kept"):
        index.rebuild(live)
    assert (live.chunks_dir / "chunks.jsonl").read_text(encoding="utf-8") == before
    assert live.manifest.read_text(encoding="utf-8") == manifest_before
    assert not (live.processed_dir.parent / ".index_build").exists()


def test_stale_index_is_refused_when_auto_rebuild_is_off(live):
    from src.index import Kora, StaleIndexError
    Kora.open(live)
    (live.raw_dir / "new.md").write_text("# New\n\nSomething new to index here.\n", encoding="utf-8")
    with pytest.raises(StaleIndexError, match="added: new.md"):
        Kora.open(live, auto_rebuild=False)


def test_empty_raw_folder_is_an_error_not_an_empty_index(tmp_path):
    from src.index import IndexBuildError, Paths, rebuild
    (tmp_path / "raw").mkdir()
    with pytest.raises(IndexBuildError, match="empty"):
        rebuild(Paths(raw_dir=tmp_path / "raw", processed_dir=tmp_path / "p", chunks_dir=tmp_path / "c",
                      intelligence_dir=tmp_path / "i", manifest=tmp_path / "m.json"))


def test_every_answer_is_logged(live):
    from src.index import Kora
    kora = Kora.open(live)
    kora.ask("How many days of annual leave do employees get?")
    records = [json.loads(l) for l in (live.log_dir / "answers.jsonl").read_text(encoding="utf-8").splitlines()]
    assert records[-1]["status"] == "answered" and records[-1]["retrieval"]["elapsed_ms"] >= 0
