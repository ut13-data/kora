"""
Stage 2: Extraction

Turns a FileInfo (from Stage 1) into usable content:
  - PDF, DOCX, Markdown, TXT  -> ExtractedDocument (plain text)
  - CSV                       -> ExtractedTable (list of row dicts)

This module does NOT clean the text (Stage 3) or chunk it (Stage 5).
It only pulls raw content out of the file format.
"""

import csv
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from src.ingestion.file_intake import FileInfo, FileType


class ExtractionError(Exception):
    """Raised when a file's content can't be read or parsed."""
    pass


@dataclass
class ExtractedDocument:
    """Normalized text content for PDF, DOCX, Markdown, and TXT files."""
    file_info: FileInfo
    text: str


@dataclass
class ExtractedTable:
    """Structured content for CSV files. Each row is a dict keyed by column header."""
    file_info: FileInfo
    rows: list[dict] = field(default_factory=list)


def _extract_txt(file_info: FileInfo) -> ExtractedDocument:
    text = Path(file_info.source_path).read_text(encoding="utf-8")
    return ExtractedDocument(file_info=file_info, text=text)


def _extract_markdown(file_info: FileInfo) -> ExtractedDocument:
    # Same approach as TXT for now — we keep the raw Markdown syntax
    # (headers, etc.) intact. Stage 3 (Cleaning) decides what to do with it.
    text = Path(file_info.source_path).read_text(encoding="utf-8")
    return ExtractedDocument(file_info=file_info, text=text)


def _extract_docx(file_info: FileInfo) -> ExtractedDocument:
    doc = DocxDocument(file_info.source_path)
    paragraphs = [p.text for p in doc.paragraphs]
    text = "\n".join(paragraphs)
    return ExtractedDocument(file_info=file_info, text=text)


def _extract_pdf(file_info: FileInfo) -> ExtractedDocument:
    reader = PdfReader(file_info.source_path)
    page_texts = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(page_texts)
    return ExtractedDocument(file_info=file_info, text=text)


def _extract_csv(file_info: FileInfo) -> ExtractedTable:
    rows = []
    with open(file_info.source_path, encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(dict(row))
    return ExtractedTable(file_info=file_info, rows=rows)


# Maps each FileType to the function that knows how to extract it.
_EXTRACTORS = {
    FileType.TXT: _extract_txt,
    FileType.MARKDOWN: _extract_markdown,
    FileType.DOCX: _extract_docx,
    FileType.PDF: _extract_pdf,
    FileType.CSV: _extract_csv,
}


def extract_file(file_info: FileInfo) -> ExtractedDocument | ExtractedTable:
    """Extract content from a file, based on its FileType.

    Returns an ExtractedDocument for text-based files, or an
    ExtractedTable for CSV files.

    Raises:
        ExtractionError: if the underlying library fails to parse the file.
    """
    extractor = _EXTRACTORS.get(file_info.file_type)
    if extractor is None:
        # Shouldn't happen if Stage 1 already validated the type,
        # but we guard anyway — never trust silently.
        raise ExtractionError(f"No extractor for file type: {file_info.file_type}")

    try:
        return extractor(file_info)
    except Exception as e:
        raise ExtractionError(f"Failed to extract '{file_info.filename}': {e}") from e